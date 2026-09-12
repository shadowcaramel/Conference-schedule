# -*- coding: utf-8 -*-
"""Match programme titles to the book-of-abstracts TOC and propose Unicode notation.

Usage:
    python tools/match_abstract_titles.py           # write data/title-notation-proposals.md
    python tools/match_abstract_titles.py --apply   # write Unicode Название into programme.xlsx

Reads data/programme.xlsx and fizika-yadra-2026_extract.pdf.
Report mode does not edit the workbook.
"""
from __future__ import annotations

import re
import sys
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    DATA_DIR,
    PROGRAMME_XLSX,
    ROOT,
    SHEET_POSTERS,
    SHEET_TALKS,
    clean,
    utf8_stdout,
)

utf8_stdout()

EXTRACT_PDF = ROOT / "fizika-yadra-2026_extract.pdf"
OUT_MD = DATA_DIR / "title-notation-proposals.md"

CYR = re.compile(r"[\u0400-\u04FF]")
TR = str.maketrans({
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
    "з": "z", "и": "i", "й": "i", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
    "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c",
    "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
    "я": "ya",
})
SUP = str.maketrans("0123456789+-=()", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾")
SUP_INV = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ₙ⸴ᵐ", "0123456789+-=()n,m")

ELEMENTS = sorted(
    """H He Li Be B C N O F Ne Na Mg Al Si P S Cl Ar K Ca Sc Ti V Cr Mn Fe Co Ni Cu Zn
    Ga Ge As Se Br Kr Rb Sr Y Zr Nb Mo Tc Ru Rh Pd Ag Cd In Sn Sb Te I Xe Cs Ba La
    Ce Pr Nd Pm Sm Eu Gd Tb Dy Ho Er Tm Yb Lu Hf Ta W Re Os Ir Pt Au Hg Tl Pb Bi Po
    At Rn Fr Ra Ac Th Pa U Np Pu Am Cm Bk Cf Es Fm Md No Lr Rf Db Sg Bh Hs Mt Ds Rg
    Cn Nh Fl Mc Lv Ts Og""".split(),
    key=len,
    reverse=True,
)
EL_ALT = "|".join(re.escape(x) for x in ELEMENTS)
# Cyrillic lookalikes used in Russian titles (Ве, Не, Тс).
NUCLIDE_EL = rf"(?:{EL_ALT}|Ве|Не|Тс|Ва|В)"
NUCLIDE_RE = re.compile(
    rf"(?<![A-Za-zА-Яа-яЁё0-9.])(\d{{1,3}}(?:\s*,\s*\d{{1,3}})*)(m)?\s*({NUCLIDE_EL})(?![a-zа-яё])"
)

SECTION_START = re.compile(
    r"^(?:Пленарные доклады|Plenary talks\b|"
    r"Секция\s+(\d+)\.|Section\s+(\d+)\.|"
    r"Стендовые доклады|Poster section\b)",
    re.I,
)
ENTRY_END = re.compile(r"(?:[.…·]{3,}|(?<![.…·])…(?![.…·]))\s*(\d{1,3})\s*$")
PAGE_ONLY = re.compile(r"^\s*\d{2,4}\s*$")
SKIP_LINE = re.compile(
    r"^(?:Contents?|Содержание)\s*$|"
    r"^(?:of nuclear reactions|of atomic nuclei|and high-energy physics|"
    r"and radioecology|technologies and facilities)\s*$",
    re.I,
)
AUTHORISH = re.compile(
    r"(?:[A-ZА-ЯЁ]\.)|(?:\bon behalf\b)|(?:от имени)|(?:от коллаборации)|"
    r"(?:for the\b)|(?:and TANGRA)|(?:эксперимент )|(?:коллабораци)",
    re.I,
)


def nrm(s: str) -> str:
    s = clean(s).translate(SUP_INV).lower().replace("ё", "е")
    s = s.replace("ψ", "psi").replace("ν", "nu").replace("α", "alpha").replace("β", "beta")
    s = s.replace("γ", "gamma").replace("π", "pi").replace("δ", "delta").replace("λ", "lambda")
    s = s.replace("√", "sqrt")
    s = re.sub(r"[^a-zа-я0-9\s-]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def lat(s: str) -> str:
    s = nrm(s).translate(TR)
    return s.replace("zh", "j").replace("kh", "h")


def sim(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def last_keys(s: str) -> set[str]:
    k = lat(s)
    out = {k} if k else set()
    if k.startswith("dj") and len(k) > 3:
        out.add(k[1:])
    if k.startswith("j") and not k.startswith("dj"):
        out.add("d" + k)
    return out


def superscript_masses(num: str) -> str:
    # Ordinary "," sits on the baseline beside superscript digits. U+2E34 RAISED
    # COMMA matches the book (⁴⁰⸴ ⁴⁸Ca), not ⁴⁰,⁴⁸Ca.
    raised = re.sub(r"\s*,\s*", "\u2e34 ", num)
    return raised.translate(SUP)


def _nuclide_sub(m: re.Match) -> str:
    masses, isomer, el = m.group(1), m.group(2) or "", m.group(3)
    rest = m.string[m.end():]
    if el == "Ge" and rest.startswith("V"):
        return m.group(0)
    # Metastable m stays a normal letter at the same size as the element
    # (⁹⁹mTc), not modifier ᵐ which sits lower than the mass digits.
    return superscript_masses(masses) + isomer + el


def notation_fix(title: str) -> str:
    """Keep wording; upgrade flattened math to Unicode."""
    s = clean(title)
    s = s.replace("\uf06e", "ν").replace("", "ν")
    s = s.replace("ᵐ", "m")
    s = re.sub(r"\$J\\?/\\psi\$", "J/ψ", s, flags=re.I)
    s = re.sub(r"\$J\\?/ψ\$", "J/ψ", s, flags=re.I)
    s = re.sub(r"J\\psi", "J/ψ", s)
    s = re.sub(r"\bNUGEN\b", "νGeN", s, flags=re.I)
    s = re.sub(r"\bnuGeN\b", "νGeN", s)
    s = re.sub(r"ν\s*GeN", "νGeN", s)
    s = re.sub(r"[√]\s*[sS][\s_]*N[\s_]*N", "√sₙₙ", s)
    s = re.sub(r"\balpha\+", "α+", s)
    s = re.sub(r"K\*\s*-\s*\(892\)", "K*⁻(892)", s)
    s = re.sub(r"K\*\s*[–—−]\s*\(892\)", "K*⁻(892)", s)
    s = NUCLIDE_RE.sub(_nuclide_sub, s)
    return re.sub(r"\s+", " ", s).strip()


def person_last(tokens: list[str]) -> str:
    toks = [t.strip(" ,;") for t in tokens if t.strip(" ,;")]
    if not toks:
        return ""

    def is_init(t: str) -> bool:
        return bool(re.fullmatch(r"[A-ZА-ЯЁ]\.?", t))

    longs = [t.strip(".") for t in toks if not is_init(t) and len(t.strip(".")) > 1]
    if not longs:
        return toks[-1].strip(".")
    if is_init(toks[0]) and longs:
        return longs[-1]
    return longs[0]


def author_lasts(blob: str) -> list[str]:
    blob = re.sub(r"\([^)]*\)", " ", blob)
    people = re.split(r",|;", blob)
    out = []
    for p in people:
        toks = re.split(r"\s+", p.strip())
        last = person_last(toks)
        if last:
            out.append(last)
    return out


def looks_like_authors(line: str) -> bool:
    line = line.strip()
    if not line or PAGE_ONLY.match(line) or SKIP_LINE.match(line):
        return False
    if AUTHORISH.search(line):
        return True
    parts = [p for p in re.split(r"\s+", line) if p]
    if 2 <= len(parts) <= 6 and all(re.match(r"[A-ZА-ЯЁ]", p) for p in parts[:2]):
        if not re.search(r"\b(in|of|for|and|the|on|with|at)\b", line, re.I):
            return "," in line or len(parts) <= 4
    return False


def _section_of(line: str) -> str | None:
    m = SECTION_START.match(line.strip())
    if not m:
        return None
    if re.match(r"Пленарные|Plenary", line, re.I):
        return "P"
    if re.match(r"Стендовые|Poster", line, re.I):
        return "POST"
    return m.group(1) or m.group(2)


def _split_auth_title(kept: list[str]) -> tuple[str, str]:
    auth_lines: list[str] = []
    title_lines: list[str] = []
    taking_auth = True
    for ln in kept:
        if taking_auth and looks_like_authors(ln):
            auth_lines.append(ln)
        else:
            taking_auth = False
            title_lines.append(ln)
    if not title_lines and auth_lines:
        title_lines = auth_lines[1:]
        auth_lines = auth_lines[:1]
    return " ".join(auth_lines), clean(" ".join(title_lines))


def parse_toc(pdf_path: Path) -> list[dict]:
    try:
        from pypdf import PdfReader
    except ImportError:
        sys.stderr.write("pypdf is required: pip install pypdf\n")
        raise SystemExit(1)

    reader = PdfReader(str(pdf_path))
    lines: list[str] = []
    for page in reader.pages:
        chunk = page.extract_text() or ""
        for raw in chunk.splitlines():
            line = raw.replace("\u00a0", " ").strip()
            if not line or PAGE_ONLY.match(line):
                continue
            if re.fullmatch(r"2[5-7]\d", line):
                continue
            lines.append(line)

    section = "P"
    entries: list[dict] = []
    i = 0
    buf: list[str] = []

    def flush(page: str) -> None:
        nonlocal buf, section
        kept = []
        for ln in buf:
            sec = _section_of(ln)
            if sec:
                section = sec
                continue
            if SKIP_LINE.match(ln) or ln.lower().startswith("section "):
                continue
            kept.append(ln)
        buf = []
        if not kept or not page:
            return
        try:
            n = int(page)
        except ValueError:
            return
        if n < 3 or n > 249:
            return
        authors, title = _split_auth_title(kept)
        if not title:
            return
        entries.append({
            "section": section,
            "page": page,
            "authors": authors,
            "lasts": author_lasts(authors),
            "title": title,
        })

    while i < len(lines):
        cur = lines[i]
        sec = _section_of(cur)
        if sec and not buf:
            section = sec
            i += 1
            continue
        em = ENTRY_END.search(cur)
        if em:
            left = cur[:em.start()].rstrip(" .")
            right = cur[em.end():].strip()
            if left:
                buf.append(left)
            flush(em.group(1))
            i += 1
            if right:
                lines.insert(i, right)
            continue
        buf.append(cur)
        i += 1
    return entries


def header_map(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {clean(c): i + 1 for i, c in enumerate(row) if clean(c)}


def read_people() -> list[dict]:
    wb = openpyxl.load_workbook(PROGRAMME_XLSX, data_only=True)
    people = []
    for sheet, kind in ((SHEET_TALKS, "talk"), (SHEET_POSTERS, "poster")):
        ws = wb[sheet]
        h = header_map(ws)
        for r in range(2, ws.max_row + 1):
            pid = clean(ws.cell(r, h["ID"]).value)
            if not pid:
                continue
            sec = clean(ws.cell(r, h["Секция"]).value)
            people.append({
                "id": pid,
                "kind": kind,
                "section": "POST" if kind == "poster" else (sec or ""),
                "last": clean(ws.cell(r, h["Фамилия"]).value),
                "first": clean(ws.cell(r, h["Имя"]).value),
                "title": clean(ws.cell(r, h["Название"]).value),
            })
    wb.close()
    return people


def index_toc(toc: list[dict]) -> dict[str, list[dict]]:
    by_last: dict[str, list[dict]] = defaultdict(list)
    for e in toc:
        for ln in e["lasts"]:
            for k in last_keys(ln):
                by_last[k].append(e)
    return by_last


def match_one(p: dict, toc: list[dict], by_last: dict[str, list[dict]]) -> tuple[dict | None, str, float]:
    cands: list[dict] = []
    seen: set[int] = set()
    for k in last_keys(p["last"]):
        for e in by_last.get(k, []):
            i = id(e)
            if i not in seen:
                seen.add(i)
                cands.append(e)
    if not cands:
        pl = lat(p["last"])
        for k, rows in by_last.items():
            if sim(pl, k) >= 0.88:
                for e in rows:
                    i = id(e)
                    if i not in seen:
                        seen.add(i)
                        cands.append(e)
    if not cands:
        # last-name miss (speaker is not first author): unique high title score
        scored = sorted(((sim(nrm(p["title"]), nrm(e["title"])), e) for e in toc), key=lambda x: -x[0])
        if scored and scored[0][0] >= 0.78:
            best, row = scored[0]
            second = scored[1][0] if len(scored) > 1 else 0.0
            if best - second >= 0.08:
                return row, "title", best
        return None, "unmatched", 0.0

    def score(e: dict) -> float:
        base = sim(nrm(p["title"]), nrm(e["title"]))
        want = p["section"]
        got = e["section"]
        if want == "POST" and got == "POST":
            base += 0.06
        elif want == "P" and got == "P":
            base += 0.06
        elif want and got and want == got:
            base += 0.06
        return base

    ranked = sorted(((score(e), e) for e in cands), key=lambda x: -x[0])
    best, row = ranked[0]
    second = ranked[1][0] if len(ranked) > 1 else 0.0
    if best < 0.42:
        scored = sorted(((sim(nrm(p["title"]), nrm(e["title"])), e) for e in toc), key=lambda x: -x[0])
        if scored and scored[0][0] >= 0.78:
            tbest, trow = scored[0]
            tsecond = scored[1][0] if len(scored) > 1 else 0.0
            if tbest - tsecond >= 0.08:
                return trow, "title", tbest
        return None, "weak-title", best
    if len(ranked) > 1 and best - second < 0.04 and best < 0.78:
        return row, "ambiguous", best
    return row, "ok", best


def aligned_book(prog: str, book: str) -> str:
    """If the TOC blob swallowed the next talk, keep the matching prefix."""
    a, b = nrm(prog), nrm(book)
    if a and b.startswith(a) and len(b) > len(a) + 12:
        return prog
    return book


def md_cell(s: str) -> str:
    return (s or "").replace("|", "\\|").replace("\n", " ")
    return (s or "").replace("|", "\\|").replace("\n", " ")


def classify(p: dict, hit: dict | None, why: str, score: float) -> str:
    proposed = notation_fix(p["title"])
    if hit is None:
        return "D"
    book = aligned_book(p["title"], hit["title"])
    wording = sim(nrm(p["title"]), nrm(book))
    if wording < 0.62:
        return "B"
    if proposed != p["title"]:
        return "A"
    return "C"


def write_report(people: list[dict], toc: list[dict], matches: list[tuple]) -> None:
    buckets = {"A": [], "B": [], "C": [], "D": []}
    used_toc: set[int] = set()
    for p, hit, why, score in matches:
        buckets[classify(p, hit, why, score)].append((p, hit, why, score))
        if hit is not None:
            used_toc.add(id(hit))
    unmatched_toc = [e for e in toc if id(e) not in used_toc]

    lines = [
        "# Title notation proposals (ЯДРО-2026)",
        "",
        "Generated by `python tools/match_abstract_titles.py`. **Not applied** to `programme.xlsx`.",
        "Programme wording/language is kept; only math notation is rewritten as Unicode.",
        "",
        f"- Programme rows: **{len(people)}**",
        f"- TOC entries parsed: **{len(toc)}**",
        f"- A notation upgrade: **{len(buckets['A'])}**",
        f"- B wording/language mismatch: **{len(buckets['B'])}**",
        f"- C already fine: **{len(buckets['C'])}**",
        f"- D unmatched programme: **{len(buckets['D'])}**",
        f"- TOC leftovers (not matched to a programme row): **{len(unmatched_toc)}**",
        "",
        "Onest does not include Greek, √, or Unicode super/subscripts; those glyphs use the system fallback.",
        "Mass numbers are still proposed as Unicode superscripts so Excel/ICS stay plain text.",
        "A normal comma next to those digits sits on the baseline (not a Markdown bug); isotope lists use a raised comma: ⁴⁰⸴ ⁴⁸Ca.",
        "Isomer m stays a normal letter after the mass (⁹⁹mTc), matching the book — not a second-level subscript.",
        "",
    ]

    def block(title: str, key: str, extra: str) -> None:
        rows = buckets[key]
        lines.append(f"## {title}")
        lines.append("")
        lines.append(extra)
        lines.append("")
        if not rows:
            lines.append("_None._")
            lines.append("")
            return
        lines.append("| ID | Speaker | Score | Current | Book | Proposed |")
        lines.append("|---|---|---|---|---|---|")
        for p, hit, why, score in rows:
            book = hit["title"] if hit else "—"
            if hit:
                trimmed = aligned_book(p["title"], hit["title"])
                if trimmed != hit["title"]:
                    book = trimmed + " _(trimmed; TOC glued next talk)_"
                if hit["section"] != p["section"] and p["section"] not in ("", hit["section"]):
                    book += f" _(TOC {hit['section']}, p.{hit['page']})_"
                else:
                    book += f" _(p.{hit['page']})_"
            lines.append(
                f"| `{p['id']}` | {md_cell(p['last'])} | {score:.2f} {why} | "
                f"{md_cell(p['title'])} | {md_cell(book)} | {md_cell(notation_fix(p['title']))} |"
            )
        lines.append("")

    block(
        "A. Notation upgrade",
        "A",
        "Apply these `Название` replacements (same words, Unicode math).",
    )
    block(
        "B. Wording / language mismatch",
        "B",
        "Same speaker, different title text. Do not auto-replace with the book title. Proposed column is notation-only on the **programme** wording.",
    )
    block(
        "C. Already fine",
        "C",
        "Matched; no Unicode change suggested.",
    )
    block(
        "D. Unmatched programme",
        "D",
        "No confident TOC hit. **Proposed** is still a notation pass on the programme title (apply those Unicode bits even without a book match). Jubilee/sponsor talks are expected here.",
    )

    lines.append("## TOC leftovers")
    lines.append("")
    if not unmatched_toc:
        lines.append("_None._")
        lines.append("")
    else:
        lines.append("| TOC section | Page | Authors | Title |")
        lines.append("|---|---|---|---|")
        for e in unmatched_toc:
            lines.append(
                f"| {e['section']} | {e['page']} | {md_cell(e['authors'])} | {md_cell(e['title'])} |"
            )
        lines.append("")

    OUT_MD.write_text("\n".join(lines), encoding="utf-8")


def apply_to_xlsx() -> tuple[int, Path]:
    """Write notation_fix() into Название. Returns (changed rows, path written)."""
    wb = openpyxl.load_workbook(PROGRAMME_XLSX)
    n = 0
    samples = []
    for sheet in (SHEET_TALKS, SHEET_POSTERS):
        ws = wb[sheet]
        h = header_map(ws)
        for r in range(2, ws.max_row + 1):
            if not clean(ws.cell(r, h["ID"]).value):
                continue
            cell = ws.cell(r, h["Название"])
            old = clean(cell.value)
            if not old:
                continue
            new = notation_fix(old)
            if new != old:
                cell.value = new
                n += 1
                if len(samples) < 8:
                    pid = clean(ws.cell(r, h["ID"]).value)
                    samples.append(f"  {pid}: {old} → {new}")
    dest = PROGRAMME_XLSX
    try:
        wb.save(dest)
    except PermissionError:
        dest = DATA_DIR / "_programme_update.xlsx"
        wb.save(dest)
    for line in samples:
        print(line)
    return n, dest


def main() -> int:
    apply = "--apply" in sys.argv
    if not PROGRAMME_XLSX.exists():
        print(f"missing {PROGRAMME_XLSX}")
        return 1
    if not EXTRACT_PDF.exists():
        print(f"missing {EXTRACT_PDF}")
        return 1

    if apply:
        try:
            n, dest = apply_to_xlsx()
        except PermissionError:
            print("Файл programme.xlsx открыт в Excel и заблокирован. Закройте его и повторите.")
            return 1
        print(f"updated {n} titles → {dest}")
        if dest != PROGRAMME_XLSX:
            print("Excel kept the original locked; wrote data/_programme_update.xlsx (build.py will pick it up).")

    toc = parse_toc(EXTRACT_PDF)
    people = read_people()
    by_last = index_toc(toc)
    matches = []
    for p in people:
        hit, why, score = match_one(p, toc, by_last)
        matches.append((p, hit, why, score))
    write_report(people, toc, matches)
    a = sum(1 for p, h, w, s in matches if classify(p, h, w, s) == "A")
    b = sum(1 for p, h, w, s in matches if classify(p, h, w, s) == "B")
    c = sum(1 for p, h, w, s in matches if classify(p, h, w, s) == "C")
    d = sum(1 for p, h, w, s in matches if classify(p, h, w, s) == "D")
    print(f"TOC {len(toc)}; programme {len(people)}; A {a}; B {b}; C {c}; D {d}")
    print(f"wrote {OUT_MD}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
