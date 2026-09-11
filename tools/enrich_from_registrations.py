# -*- coding: utf-8 -*-
"""Match programme speakers to the registration dump and fill names / patronyms / emails.

Usage:
    python tools/enrich_from_registrations.py --dry-run
    python tools/enrich_from_registrations.py

Edits data/programme.xlsx in place (Отчество column, mixed-script names, emails).
Does not deploy. Re-run tools/build.py after a real write.
"""
from __future__ import annotations

import argparse
import copy
import re
import sys
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import (  # noqa: E402
    PROGRAMME_XLSX,
    ROOT,
    SHEET_BLOCKS,
    SHEET_POSTERS,
    SHEET_TALKS,
    clean,
    utf8_stdout,
)

utf8_stdout()

DEFAULT_REPORT = ROOT / "Report Combi 2026-08-31 11_20_12(+0 UTG).xlsx"

CYR = re.compile(r"[\u0400-\u04FF]")
LAT = re.compile(r"[A-Za-z]")
PATR_RE = re.compile(
    r"(?:ович|евич|ёвич|овна|евна|ична|инична|ovich|evich|ovna|evna|ич|ich)$",
    re.I,
)
TR = str.maketrans({
    "а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh",
    "з": "z", "и": "i", "й": "i", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o",
    "п": "p", "р": "r", "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c",
    "ч": "ch", "ш": "sh", "щ": "sch", "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu",
    "я": "ya",
})
GIVEN_FOLD = (
    ("alexey", "aleksei"),
    ("aleksey", "aleksei"),
    ("alexei", "aleksei"),
    ("yuriy", "yuri"),
    ("yury", "yuri"),
    ("anastasiia", "anastasia"),
    ("anastasiya", "anastasia"),
)

# Confirmed overrides (see plan).
FORCE_EMAIL = {
    "S1-18": "rasulova@jinr.ru",
    "S6-02": "vkondrat@theor.jinr.ru",
    "S6-03": "vkondrat@theor.jinr.ru",
    "POST-17": "vkondrat@theor.jinr.ru",
}
TAKE_REPORT_NAMES = {"P-S2", "S1-20"}  # plus any mixed-script row
KEEP_PROG_NAMES = {"P-12", "S1-18"}
DROP_IDS = {"S2-18"}
NAME_FIXES = {
    "POST-17": ("Владимир", "Кондратьев"),  # then split Николаевич from a patched first
}


def nrm(s: str) -> str:
    s = clean(s).lower().replace("ё", "е").replace(".", " ").replace(",", " ")
    s = re.sub(r"[^a-zа-я0-9\s-]", "", s)
    return re.sub(r"\s+", " ", s).strip()


def lat(s: str) -> str:
    s = nrm(s).translate(TR)
    return s.replace("zh", "j").replace("kh", "h")


def sim(a: str, b: str) -> float:
    if not a or not b:
        return 0.0
    return SequenceMatcher(None, a, b).ratio()


def first_token(s: str) -> str:
    parts = nrm(s).replace("-", " ").split()
    return parts[0] if parts else ""


def last_keys(s: str) -> set[str]:
    k = lat(s)
    out = {k} if k else set()
    if k.startswith("dj") and len(k) > 3:
        out.add(k[1:])
    if k.startswith("j") and not k.startswith("dj"):
        out.add("d" + k)
    return out


def fold_given(s: str) -> str:
    s = lat(s)
    for src, dst in GIVEN_FOLD:
        if s.startswith(src):
            return dst + s[len(src):]
    return s


def is_initials(first: str) -> bool:
    parts = [p for p in re.split(r"[\s.]+", clean(first).replace(",", " ")) if p]
    return bool(parts) and all(len(p) <= 2 for p in parts)


def mixed_script(first: str, last: str) -> bool:
    blob = f"{first} {last}"
    return bool(CYR.search(blob) and LAT.search(blob))


def given_ok(prog_first: str, reg_first: str) -> bool:
    if is_initials(prog_first):
        return True
    a, b = fold_given(first_token(prog_first)), fold_given(first_token(reg_first))
    if not a or not b:
        return True
    if a == b:
        return True
    if a.startswith(b) or b.startswith(a):
        return len(min(a, b, key=len)) >= 3
    return sim(a, b) >= 0.8


def split_patronym(first: str) -> tuple[str, str]:
    text = clean(first).replace(",", " ")
    text = re.sub(r"\s+", " ", text).strip()
    text = text.replace("Вмкторовна", "Викторовна")
    parts = text.split()
    if len(parts) < 2:
        return text, ""
    last = parts[-1]
    if PATR_RE.search(last):
        return " ".join(parts[:-1]), last
    return text, ""


def header_map(ws) -> dict[str, int]:
    row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    return {clean(c): i + 1 for i, c in enumerate(row) if clean(c)}


def load_regs(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Worksheet"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [clean(c) for c in rows[3]]
    out = []
    for row in rows[4:]:
        rec = {hdr[i]: clean(row[i]) if i < len(row) else "" for i in range(len(hdr))}
        if rec.get("Last name") or rec.get("First Name") or rec.get("Email"):
            out.append(rec)
    wb.close()
    return out


def uniq_regs(regs: list[dict]) -> list[dict]:
    seen: set[str] = set()
    out = []
    for r in regs:
        key = r.get("Email") or f"{r.get('First Name')}|{r.get('Last name')}"
        if key in seen:
            continue
        seen.add(key)
        out.append(r)
    return out


def match_reg(person: dict, regs: list[dict], by_last: dict[str, list[dict]]) -> tuple[dict | None, str]:
    last = person["last"]
    first = person["first"]
    title = nrm(person["title"])
    cands: list[dict] = []
    for k in last_keys(last):
        cands.extend(by_last.get(k, []))
        cands.extend(by_last.get("swapped:" + k, []))
    if not cands:
        pl = lat(last)
        seen_k: set[str] = set()
        for k, rs in by_last.items():
            if k.startswith("swapped:") or k in seen_k:
                continue
            if sim(pl, k) >= 0.86:
                seen_k.add(k)
                cands.extend(rs)
    cands = uniq_regs(cands)
    filtered = [r for r in cands if given_ok(first, r.get("First Name") or "")]
    pool = filtered if filtered else ([] if cands else [])

    def by_title(rows: list[dict]) -> tuple[dict | None, str]:
        if not rows or not title:
            return None, ""
        scored = sorted(((sim(title, nrm(r.get("Title") or "")), r) for r in rows), key=lambda x: -x[0])
        best, row = scored[0]
        second = scored[1][0] if len(scored) > 1 else 0.0
        if best >= 0.55 and (len(scored) == 1 or best - second >= 0.08):
            return row, f"title {best:.2f}"
        return None, ""

    if len(pool) == 1:
        return pool[0], "name"
    if len(pool) > 1:
        hit, why = by_title(pool)
        if hit:
            return hit, why
        return None, "ambiguous"
    if cands and not filtered:
        # last matched but first disagreed — do not take the unique-last false friend
        hit, why = by_title(cands)
        if hit and given_ok(first, hit.get("First Name") or ""):
            return hit, why
        hit, why = by_title(regs)
        if hit:
            return hit, why + "-global"
        return None, "first-mismatch"
    hit, why = by_title(regs)
    if hit:
        return hit, why + "-global"
    return None, "unmatched"


def index_regs(regs: list[dict]) -> dict[str, list[dict]]:
    by_last: dict[str, list[dict]] = defaultdict(list)
    for r in regs:
        for k in last_keys(r.get("Last name") or ""):
            by_last[k].append(r)
        for k in last_keys(r.get("First Name") or ""):
            by_last["swapped:" + k].append(r)
    return by_last


def read_people(ws, kind: str) -> list[dict]:
    h = header_map(ws)
    people = []
    for r in range(2, ws.max_row + 1):
        pid = clean(ws.cell(r, h["ID"]).value)
        if not pid:
            continue
        people.append({
            "row": r,
            "kind": kind,
            "id": pid,
            "last": clean(ws.cell(r, h["Фамилия"]).value),
            "first": clean(ws.cell(r, h["Имя"]).value),
            "title": clean(ws.cell(r, h["Название"]).value),
        })
    return people


def ensure_middle_column(ws, after_header: str = "Имя") -> None:
    h = header_map(ws)
    if "Отчество" in h:
        return
    col = h[after_header] + 1
    ws.insert_cols(col)
    cell = ws.cell(1, col, "Отчество")
    src = ws.cell(1, h[after_header])
    cell.font = copy.copy(src.font)
    cell.fill = copy.copy(src.fill)
    cell.alignment = copy.copy(src.alignment)
    letter = get_column_letter(col)
    ws.column_dimensions[letter].width = 18


def drop_s2_18(wb) -> None:
    ws = wb[SHEET_TALKS]
    h = header_map(ws)
    for r in range(ws.max_row, 1, -1):
        if clean(ws.cell(r, h["ID"]).value) == "S2-18":
            ws.delete_rows(r)
            print("dropped S2-18")
            break
    blocks = wb[SHEET_BLOCKS]
    bh = header_map(blocks)
    for r in range(2, blocks.max_row + 1):
        sec = clean(blocks.cell(r, bh["Секция"]).value)
        spec = clean(blocks.cell(r, bh["Доклады"]).value)
        if sec == "2" and spec == "15-21":
            blocks.cell(r, bh["Доклады"]).value = "15-17,19-21"
            print("S2 Tuesday 16:15 block: 15-21 -> 15-17,19-21")


def apply_person(person: dict, reg: dict | None, why: str) -> dict:
    pid = person["id"]
    first, last = person["first"], person["last"]
    email = ""
    take_names = False
    if reg:
        email = reg.get("Email") or ""
        mixed = mixed_script(first, last)
        take_names = (mixed or pid in TAKE_REPORT_NAMES) and pid not in KEEP_PROG_NAMES
        if take_names:
            first = reg.get("First Name") or first
            last = reg.get("Last name") or last
    if pid in FORCE_EMAIL:
        email = FORCE_EMAIL[pid]
        why = (why + "+force-email").strip("+")
    if pid == "POST-17":
        first, last = "Владимир Николаевич", "Кондратьев"
        take_names = True
        why += "+kondratiev"
    first, middle = split_patronym(first)
    return {
        "id": pid,
        "first": first,
        "middle": middle,
        "last": last,
        "email": email,
        "why": why,
        "take_names": take_names,
        "from": f"{reg.get('First Name','')} {reg.get('Last name','')}".strip() if reg else "",
    }


def run(report_path: Path, dry_run: bool) -> int:
    if not report_path.exists():
        print(f"Registration dump not found: {report_path}")
        return 1
    if not PROGRAMME_XLSX.exists():
        print(f"Programme book not found: {PROGRAMME_XLSX}")
        return 1

    regs = load_regs(report_path)
    by_last = index_regs(regs)
    print(f"registrations: {len(regs)}")

    wb = openpyxl.load_workbook(PROGRAMME_XLSX)
    if not dry_run:
        drop_s2_18(wb)
        ensure_middle_column(wb[SHEET_TALKS])
        ensure_middle_column(wb[SHEET_POSTERS])

    talks_ws = wb[SHEET_TALKS]
    posters_ws = wb[SHEET_POSTERS]
    people = read_people(talks_ws, "talk") + read_people(posters_ws, "poster")
    people = [p for p in people if p["id"] not in DROP_IDS]

    results = []
    unmatched = []
    ambiguous = []
    for p in people:
        if p["id"] in FORCE_EMAIL or p["id"] in TAKE_REPORT_NAMES or p["id"] == "POST-17":
            reg, why = match_reg(p, regs, by_last)
            if p["id"] in FORCE_EMAIL:
                want = FORCE_EMAIL[p["id"]].lower()
                forced = next((r for r in regs if (r.get("Email") or "").lower() == want), reg)
                row = apply_person(p, forced or reg, why or "force")
            else:
                row = apply_person(p, reg, why or "special")
            results.append((p, row))
            continue
        reg, why = match_reg(p, regs, by_last)
        row = apply_person(p, reg, why)
        results.append((p, row))
        if why == "unmatched" and not row["email"]:
            unmatched.append(p)
        elif why == "ambiguous" and not row["email"]:
            ambiguous.append(p)

    emails = sum(1 for _, r in results if r["email"])
    print(f"people {len(results)}; emails {emails}; unmatched {len(unmatched)}; ambiguous {len(ambiguous)}")
    print("\n=== mixed / specials ===")
    for p, r in results:
        if r["take_names"] or p["id"] in FORCE_EMAIL or mixed_script(p["first"], p["last"]):
            print(
                f"  {p['id']:8} {p['first']} {p['last']}  ->  "
                f"{r['first']} {r['middle']} {r['last']}  <{r['email']}>  [{r['why']}]"
            )
    if unmatched:
        print("\n=== unmatched ===")
        for p in unmatched:
            print(f"  {p['id']} {p['first']} {p['last']} / {p['title'][:60]}")
    if ambiguous:
        print("\n=== ambiguous ===")
        for p in ambiguous:
            print(f"  {p['id']} {p['first']} {p['last']}")

    if dry_run:
        print("\ndry-run: programme.xlsx not written")
        wb.close()
        return 0

    for p, r in results:
        ws = talks_ws if p["kind"] == "talk" else posters_ws
        h = header_map(ws)
        row = p["row"]
        # rows may have shifted after S2-18 delete — locate by ID
        found = None
        for rr in range(2, ws.max_row + 1):
            if clean(ws.cell(rr, h["ID"]).value) == p["id"]:
                found = rr
                break
        if found is None:
            print(f"skip missing row {p['id']}")
            continue
        ws.cell(found, h["Имя"]).value = r["first"] or None
        ws.cell(found, h["Отчество"]).value = r["middle"] or None
        ws.cell(found, h["Фамилия"]).value = r["last"] or None
        ws.cell(found, h["Email"]).value = r["email"] or None

    wb.save(PROGRAMME_XLSX)
    print(f"\nwrote {PROGRAMME_XLSX}")
    return 0


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = ap.parse_args(argv)
    return run(args.report, args.dry_run)


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
